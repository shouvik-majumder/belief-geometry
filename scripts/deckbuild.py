"""Minimal PowerPoint builder: white slides, one title, one figure, occasional bullets.

Deliberately plain. No theme, no colours, no decoration; the figures carry the content.
Also renders PNG mock-ups of each slide so the layout can be checked without PowerPoint.
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6
TITLE_TOP, TITLE_H = 0.35, 0.75
BOTTOM_KEEP = 0.5            # space reserved for the slide number
BLACK, GREY = RGBColor(0, 0, 0), RGBColor(0x66, 0x66, 0x66)
FONT = "Calibri"


def _textbox(slide, left, top, width, height, text, size, bold=False, colour=BLACK,
             align_centre=False, bullets=False, line_spacing=1.1):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text if isinstance(text, (list, tuple)) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        if bullets:
            p.space_after = Pt(6)
        if align_centre:
            from pptx.enum.text import PP_ALIGN
            p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = ("\u2022  " + line) if (bullets and line) else line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = colour
        run.font.name = FONT
    return box


def _fit(img_path: Path, box_w: float, box_h: float) -> tuple[float, float]:
    with Image.open(img_path) as im:
        w, h = im.size
    scale = min(box_w / w, box_h / h)
    return w * scale, h * scale


class Deck:
    """Accumulates slides and can write both the .pptx and PNG previews."""

    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]
        self.spec: list[dict] = []

    # --------------------------------------------------------------- slides
    def title_slide(self, title: str, lines: list[str]) -> None:
        slide = self.prs.slides.add_slide(self.blank)
        _textbox(slide, MARGIN, 2.5, SLIDE_W - 2 * MARGIN, 1.3, title, 36, bold=True, align_centre=True)
        _textbox(slide, MARGIN, 4.0, SLIDE_W - 2 * MARGIN, 1.6, lines, 16, colour=GREY,
                 align_centre=True, line_spacing=1.35)
        self.spec.append({"kind": "title", "title": title, "lines": lines})

    def slide(self, title: str, image: str | Path | None = None, bullets: list[str] | None = None,
              notes: str | None = None, bullet_size: int = 15) -> None:
        slide = self.prs.slides.add_slide(self.blank)
        _textbox(slide, MARGIN, TITLE_TOP, SLIDE_W - 2 * MARGIN, TITLE_H, title, 26, bold=True)

        top = TITLE_TOP + TITLE_H + 0.15
        if bullets and image is None:
            bullet_size = max(bullet_size, 18)
        if bullets:
            h = (0.34 if bullet_size <= 16 else 0.44) * len(bullets) + 0.1
            _textbox(slide, MARGIN + 0.1, top, SLIDE_W - 2 * MARGIN - 0.2, h, bullets,
                     bullet_size, bullets=True, line_spacing=1.15)
            top += h + 0.22

        img_box = None
        if image and not Path(image).exists():
            print(f"  WARNING: missing figure, slide built without it: {Path(image).name}")
            image = None
        if image:
            path = Path(image)
            avail_w = SLIDE_W - 2 * MARGIN
            avail_h = SLIDE_H - top - BOTTOM_KEEP
            w, h = _fit(path, avail_w, avail_h)
            left = (SLIDE_W - w) / 2
            slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(w), Inches(h))
            img_box = (left, top, w, h)

        n = len(self.prs.slides.__iter__.__self__._sldIdLst)  # 1-based slide number
        _textbox(slide, SLIDE_W - MARGIN - 0.6, SLIDE_H - 0.42, 0.6, 0.3, str(n), 10, colour=GREY)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        self.spec.append({"kind": "content", "title": title, "bullets": bullets or [],
                          "image": str(image) if image else None, "img_box": img_box, "number": n})

    # --------------------------------------------------------------- output
    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path

    def preview(self, out_dir: str | Path, dpi: int = 80) -> list[Path]:
        """Render each slide to PNG with Pillow, mirroring the pptx geometry, for layout QA."""
        out_dir = Path(out_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        for f in out_dir.glob("slide_*.png"):
            f.unlink()
        W, H = int(SLIDE_W * dpi), int(SLIDE_H * dpi)

        def font(size_pt: int, bold: bool = False):
            name = "calibrib.ttf" if bold else "calibri.ttf"
            try:
                return ImageFont.truetype(f"C:/Windows/Fonts/{name}", int(size_pt * dpi / 72))
            except OSError:
                return ImageFont.load_default()

        def wrap(draw, text, fnt, max_px):
            words, lines, cur = text.split(), [], ""
            for w in words:
                trial = (cur + " " + w).strip()
                if draw.textlength(trial, font=fnt) <= max_px or not cur:
                    cur = trial
                else:
                    lines.append(cur); cur = w
            if cur:
                lines.append(cur)
            return lines

        paths = []
        for i, s in enumerate(self.spec, 1):
            canvas = Image.new("RGB", (W, H), "white")
            d = ImageDraw.Draw(canvas)
            if s["kind"] == "title":
                f1, f2 = font(36, True), font(16)
                y = int(2.5 * dpi)
                for line in wrap(d, s["title"], f1, W - int(2 * MARGIN * dpi)):
                    d.text(((W - d.textlength(line, font=f1)) / 2, y), line, font=f1, fill="black")
                    y += int(f1.size * 1.25)
                y = int(4.0 * dpi)
                for line in s["lines"]:
                    d.text(((W - d.textlength(line, font=f2)) / 2, y), line, font=f2, fill="#666666")
                    y += int(f2.size * 1.5)
            else:
                ft = font(26, True)
                x0, y = int(MARGIN * dpi), int(TITLE_TOP * dpi)
                for line in wrap(d, s["title"], ft, W - int(2 * MARGIN * dpi)):
                    d.text((x0, y), line, font=ft, fill="black")
                    y += int(ft.size * 1.2)
                if s["bullets"]:
                    fb = font(18 if not s['img_box'] else 15)
                    yb = int((TITLE_TOP + TITLE_H + 0.15) * dpi)
                    for b in s["bullets"]:
                        for j, line in enumerate(wrap(d, "\u2022  " + b, fb, W - int(2.2 * MARGIN * dpi))):
                            d.text((int((MARGIN + 0.1) * dpi), yb), line, font=fb, fill="black")
                            yb += int(fb.size * 1.3)
                if s["img_box"]:
                    left, top, w, h = s["img_box"]
                    with Image.open(s["image"]) as im:
                        im = im.convert("RGB").resize((max(int(w * dpi), 1), max(int(h * dpi), 1)))
                        canvas.paste(im, (int(left * dpi), int(top * dpi)))
                d.text((W - int((MARGIN + 0.2) * dpi), H - int(0.42 * dpi)), str(s["number"]),
                       font=font(10), fill="#666666")
            p = out_dir / f"slide_{i:02d}.png"
            canvas.save(p)
            paths.append(p)
        return paths
